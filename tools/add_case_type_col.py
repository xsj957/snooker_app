# -*- coding: utf-8 -*-
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter, column_index_from_string
from datetime import datetime

src = 'iter-v1.0.1-第二期/测试用例.xlsx'
wb = openpyxl.load_workbook(src)

NEW_COL = 11  # Insert at K (after 优先级 at J=10)
COL_WIDTH_NEW = 10
FILL_HEADER = PatternFill(start_color='FF2F5496', end_color='FF2F5496', fill_type='solid')
FONT_HEADER = Font(name='微软雅黑', size=10, bold=True, color='FFFFFFFF')
FONT_TC = Font(name='微软雅黑', size=10)
FILL_TC = PatternFill(start_color='FFFCE4EC', end_color='FFFCE4EC', fill_type='solid')
THIN = Side(style='thin')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
ALIGN_CENTER = Alignment(horizontal='center', vertical='center', wrapText=True)
ALIGN_LEFT = Alignment(horizontal='left', vertical='center', wrapText=True)

def parse_cell(cell_str):
    """Parse cell reference like 'A2' into (col_letter, row_num)."""
    import re
    m = re.match(r'([A-Z]+)(\d+)', cell_str)
    return m.group(1), int(m.group(2))

def shift_merge_range(mr_str, insert_col):
    """Shift a merged cell range if it's at or after the insert column."""
    parts = mr_str.split(':')
    if len(parts) != 2:
        return mr_str
    min_cell, max_cell = parts
    min_col_letter, min_row = parse_cell(min_cell)
    max_col_letter, max_row = parse_cell(max_cell)
    min_col_idx = column_index_from_string(min_col_letter)
    max_col_idx = column_index_from_string(max_col_letter)

    # If merge range starts at or after insert column, shift entire range
    if min_col_idx >= insert_col:
        new_min = f'{get_column_letter(min_col_idx + 1)}{min_row}'
        new_max = f'{get_column_letter(max_col_idx + 1)}{max_row}'
        return f'{new_min}:{new_max}'
    # If merge range spans the insert column, expand the max end
    elif max_col_idx >= insert_col:
        new_max = f'{get_column_letter(max_col_idx + 1)}{max_row}'
        return f'{mr_str.split(":")[0]}:{new_max}'
    # Otherwise unchanged
    return mr_str

for sn in wb.sheetnames:
    ws = wb[sn]
    print(f'\n=== {sn} ===')

    # 1. Save merged cell info
    old_merges = [str(mr) for mr in ws.merged_cells.ranges]
    print(f'  合并范围: {old_merges}')

    # 2. Unmerge all merged cells before insert
    for mr in old_merges:
        ws.unmerge_cells(mr)

    # 3. Insert column at position 11
    ws.insert_cols(NEW_COL)

    # 4. Set column width
    ws.column_dimensions[get_column_letter(NEW_COL)].width = COL_WIDTH_NEW

    # 5. Shift merged cell ranges and re-merge
    for mr in old_merges:
        new_mr = shift_merge_range(mr, NEW_COL)
        if new_mr != mr:
            print(f'  {mr} → {new_mr}')
        ws.merge_cells(new_mr)

    # 6. Style header row for new column
    header_cell = ws.cell(row=1, column=NEW_COL)
    header_cell.value = '用例类型'
    header_cell.font = FONT_HEADER
    header_cell.fill = FILL_HEADER
    header_cell.border = BORDER
    header_cell.alignment = ALIGN_CENTER

    # 7. Style TC rows and section headers in new column
    for r in range(2, ws.max_row + 1):
        a = str(ws.cell(row=r, column=1).value or '').strip()
        b = str(ws.cell(row=r, column=2).value or '').strip()

        cell = ws.cell(row=r, column=NEW_COL)

        if a and not b and not a.startswith('斯诺克大师V') and not a.startswith('SM-') and r > 1:
            # Section header
            cell.fill = FILL_HEADER
            cell.font = FONT_HEADER
            cell.border = BORDER
            cell.alignment = ALIGN_LEFT
        elif a.startswith('斯诺克大师V1.0.1-') or a.startswith('SM-'):
            # TC row
            cell.fill = FILL_TC
            cell.font = FONT_TC
            cell.border = BORDER
            cell.alignment = ALIGN_CENTER

    # 8. Add data validation dropdown for new column
    dv = DataValidation(type='list', formula1='"回归,普通"', allow_blank=True)
    dv.error = '请选择"回归"或"普通"'
    dv.errorTitle = '无效输入'
    dv.prompt = '请选择用例类型'
    dv.promptTitle = '用例类型'
    col_letter = get_column_letter(NEW_COL)
    dv.sqref = f'{col_letter}2:{col_letter}10000'
    ws.add_data_validation(dv)
    print(f'  下拉验证: {col_letter}2:{col_letter}10000 → 回归,普通')

    # 9. Verify new header
    h = ws.cell(row=1, column=NEW_COL).value
    print(f'  Col{NEW_COL} header = {h}')

    # Show first few TC rows in new column
    for r in range(2, 5):
        a = ws.cell(row=r, column=1).value
        v = ws.cell(row=r, column=NEW_COL).value
        print(f'  R{r} {a}: 用例类型={v}')

    # Check merged cells
    merges = [str(mr) for mr in ws.merged_cells.ranges]
    print(f'  最终合并: {merges[:3]}')

# Save
ts = datetime.now().strftime('%Y%m%d_%H%M')
dst = f'iter-v1.0.1-第二期/测试用例-{ts}.xlsx'
wb.save(dst)
print(f'\n已保存: {dst}')

# Final verification
print(f'\n{"="*50}')
wb2 = openpyxl.load_workbook(dst)
for sn in wb2.sheetnames:
    ws = wb2[sn]
    h = ws.cell(row=1, column=11).value
    dv_count = len(ws.data_validations.dataValidation)
    print(f'{sn}: Col11="{h}", 下拉验证{dv_count}个, TC行数={sum(1 for r in range(2, ws.max_row+1) if str(ws.cell(row=r, column=1).value or "").startswith("斯诺克大师"))}')
