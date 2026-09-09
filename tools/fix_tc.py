# -*- coding: utf-8 -*-
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl
from copy import copy
from datetime import datetime

src = 'iter-v1.0.1-第二期/测试用例-20260909_2308.xlsx'
timestamp = datetime.now().strftime('%Y%m%d_%H%M')
dst = f'iter-v1.0.1-第二期/测试用例-{timestamp}.xlsx'

wb = openpyxl.load_workbook(src)

def insert_tc_at(ws, target_row, tc_data, style_row=None):
    ws.insert_rows(target_row)
    if style_row is None:
        style_row = target_row + 1
    col_map = {'A':1,'B':2,'C':3,'D':4,'E':5,'F':6,'G':7,'H':8,'I':9,'J':10}
    for col_letter, col_idx in col_map.items():
        cell = ws.cell(row=target_row, column=col_idx)
        cell.value = tc_data.get(col_letter, '')
        ref = ws.cell(row=style_row, column=col_idx)
        cell.font = copy(ref.font)
        cell.fill = copy(ref.fill)
        cell.border = copy(ref.border)
        cell.alignment = copy(ref.alignment)

def renumber_tcs(ws, start_row, end_row):
    prefix = '斯诺克大师V1.0.1-'
    counter = 1
    for r in range(start_row, end_row + 1):
        val = ws.cell(row=r, column=1).value
        if val and str(val).startswith(prefix):
            try:
                num = int(str(val).split('-')[-1])
                counter = num
                break
            except:
                pass
    for r in range(start_row, end_row + 1):
        val = ws.cell(row=r, column=1).value
        if val and str(val).startswith(prefix):
            ws.cell(row=r, column=1).value = f'{prefix}{counter:04d}'
            counter += 1

# Find target rows in 安卓
ws = wb['安卓']
rows = {}
for r in range(1, ws.max_row + 1):
    val = str(ws.cell(row=r, column=1).value or '')
    for key in ['0021', '0066', '0067']:
        if f'-{key}' in val:
            rows[key] = r

print(f'安卓 目标行: {rows}')

# Insert from bottom to top
# 1. After 0067: 角标边界条件
if '0067' in rows:
    insert_tc_at(ws, rows['0067'] + 1, {
        'A': 'PLACEHOLDER',
        'B': '斯诺克大师国内APP',
        'C': '个人数据模块',
        'D': '功能',
        'E': '我的视频new角标-边界条件',
        'F': '首页我的视频角标仅在视频数>0且未查看数>1时才显示',
        'G': '用户账号下有已制作完成的视频',
        'H': '1.首页查看我的视频角标显示条件\n2.当仅有1个未查看视频时观察角标\n3.当有2个及以上未查看视频时观察角标\n4.当所有视频都已查看时观察角标',
        'I': '1.仅1个未查看视频时不显示角标\n2.>=2个未查看视频时角标显示对应数量\n3.全部已查看时角标消失\n4.视频数为0时角标不显示',
        'J': 'P1',
    }, style_row=rows['0067'])
    print(f'  安卓: R{rows["0067"]+1} 新增角标边界条件')

# 2. After 0066: 视频数量分场景
if '0066' in rows:
    insert_tc_at(ws, rows['0066'] + 1, {
        'A': 'PLACEHOLDER',
        'B': '斯诺克大师国内APP',
        'C': '个人数据模块',
        'D': '功能',
        'E': '我的视频数量-计入与不计入场景',
        'F': '我的视频数量按新规则分别验证工控机完成/app完成有缓存/腾讯云过期/制作中/缓存删除等场景计数正确',
        'G': '用户账号下存在各类状态的视频',
        'H': '1.工控机制作完成且腾讯云未过期的视频观察数量\n2.app制作完成有本地缓存的视频(腾讯云已过期)观察数量\n3.工控机制作但腾讯云已过期的视频观察数量\n4.app制作中(待下载/下载中)的视频观察数量\n5.app本地制作完成但缓存被删除的视频观察数量',
        'I': '1.工控机完成且未过期视频计入总数\n2.app有本地缓存视频计入总数(不论腾讯云是否过期)\n3.腾讯云已过期视频不计入\n4.制作中视频不计入\n5.缓存被删除的视频不计入',
        'J': 'P1',
    }, style_row=rows['0066'])
    print(f'  安卓: R{rows["0066"]+1} 新增视频数量分场景')

# 3. Modify 0021 to include 120分
if '0021' in rows:
    r = rows['0021']
    old_h = str(ws.cell(row=r, column=8).value or '')
    old_f = str(ws.cell(row=r, column=6).value or '')
    if '120' not in old_h:
        ws.cell(row=r, column=8).value = old_h.replace('140分', '120/130/140分')
        ws.cell(row=r, column=6).value = old_f.replace('140分', '120/130/140分')
        print(f'  安卓: 0021 增加120分边界值')

# Find target rows in iOS
ws_ios = wb['国内iOS']
rows_ios = {}
for r in range(1, ws_ios.max_row + 1):
    val = str(ws_ios.cell(row=r, column=1).value or '')
    for key in ['0059', '0060']:
        if f'-{key}' in val:
            rows_ios[key] = r

print(f'\niOS 目标行: {rows_ios}')

# Insert from bottom to top
# 1. After 0060: 角标边界条件
if '0060' in rows_ios:
    insert_tc_at(ws_ios, rows_ios['0060'] + 1, {
        'A': 'PLACEHOLDER',
        'B': '斯诺克大师国内APP',
        'C': '个人数据模块',
        'D': '功能',
        'E': '我的视频new角标-边界条件',
        'F': '首页我的视频角标仅在视频数>0且未查看数>1时才显示',
        'G': '用户账号下(iOS)有已制作完成的视频',
        'H': '1.首页查看我的视频角标显示条件\n2.当仅有1个未查看视频时观察角标\n3.当有2个及以上未查看视频时观察角标\n4.当所有视频都已查看时观察角标',
        'I': '1.仅1个未查看视频时不显示角标\n2.>=2个未查看视频时角标显示对应数量\n3.全部已查看时角标消失\n4.视频数为0时角标不显示',
        'J': 'P1',
    }, style_row=rows_ios['0060'])
    print(f'  iOS: R{rows_ios["0060"]+1} 新增角标边界条件')

# 2. After 0059: 视频数量分场景
if '0059' in rows_ios:
    insert_tc_at(ws_ios, rows_ios['0059'] + 1, {
        'A': 'PLACEHOLDER',
        'B': '斯诺克大师国内APP',
        'C': '个人数据模块',
        'D': '功能',
        'E': '我的视频数量-计入与不计入场景',
        'F': '我的视频数量按新规则分别验证工控机完成/app完成有缓存/腾讯云过期/制作中/缓存删除等场景计数正确',
        'G': '用户账号下(iOS)存在各类状态的视频',
        'H': '1.工控机制作完成且腾讯云未过期的视频观察数量\n2.app制作完成有本地缓存的视频(腾讯云已过期)观察数量\n3.工控机制作但腾讯云已过期的视频观察数量\n4.app制作中(待下载/下载中)的视频观察数量\n5.app本地制作完成但缓存被删除的视频观察数量',
        'I': '1.工控机完成且未过期视频计入总数\n2.app有本地缓存视频计入总数(不论腾讯云是否过期)\n3.腾讯云已过期视频不计入\n4.制作中视频不计入\n5.缓存被删除的视频不计入',
        'J': 'P1',
    }, style_row=rows_ios['0059'])
    print(f'  iOS: R{rows_ios["0059"]+1} 新增视频数量分场景')

# Renumber all TCs in each sheet
print('\n=== 重新编号 ===')
for sheet_name in ['安卓', '国内iOS']:
    ws = wb[sheet_name]
    first_tc = None
    for r in range(2, ws.max_row + 1):
        val = str(ws.cell(row=r, column=1).value or '')
        if val.startswith('斯诺克大师V1.0.1-'):
            first_tc = r
            break
    if first_tc:
        renumber_tcs(ws, first_tc, ws.max_row)
        print(f'  {sheet_name}: R{first_tc}~R{ws.max_row} 编号完成')

# Save
try:
    wb.save(dst)
    print(f'\n已保存: {dst}')
except PermissionError:
    timestamp2 = datetime.now().strftime('%Y%m%d_%H%M%S')
    dst = f'iter-v1.0.1-第二期/测试用例-{timestamp2}.xlsx'
    wb.save(dst)
    print(f'\nWPS锁定，已保存为: {dst}')
